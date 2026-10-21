"""
导出 '会话列表及会话盒子.xlsx' 已实现 (50) sheet 的当前状态汇总：
- 已通过
- 跳过
- 待修复（仍失败 / 还没动）

按 row 排序，输出 markdown 报告。
"""
import re
from openpyxl import load_workbook
from pathlib import Path

EXCEL = r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx"
SHEET = "已实现 (50)"

wb = load_workbook(EXCEL, data_only=True)
ws = wb[SHEET]

cases = []
for r in range(2, ws.max_row + 1):
    text = ws.cell(r, 7).value or ""
    if not text:
        continue
    m = re.search(r'(test_[^\s\(\)（）]+\.py)', text)
    test_file = m.group(1) if m else "?"
    if text.startswith("已通过"):
        status = "已通过"
        date = re.search(r'\d{4}-\d{2}-\d{2}', text).group(0) if re.search(r'\d{4}-\d{2}-\d{2}', text) else ""
    elif text.startswith("跳过"):
        status = "跳过"
        # 抓 reason: "跳过 2026-05-21 - <reason> - 已实现..."
        rm = re.search(r'^跳过 \d{4}-\d{2}-\d{2} - ([^-]+?) - 已实现', text)
        date = re.search(r'\d{4}-\d{2}-\d{2}', text).group(0) if re.search(r'\d{4}-\d{2}-\d{2}', text) else ""
    elif text.startswith("失败"):
        status = "待修复(已尝试失败)"
        date = ""
    else:
        status = "待修复(未尝试)"
        date = ""
    cases.append({
        'row': r,
        'status': status,
        'date': date,
        'test_file': test_file,
        '功能点': ws.cell(r, 1).value or "",
        '测试点': ws.cell(r, 3).value or "",
        'progress': text,
    })

# 分组
passed = [c for c in cases if c['status'] == "已通过"]
skipped = [c for c in cases if c['status'] == "跳过"]
todo = [c for c in cases if c['status'].startswith("待修复")]

out = []
out.append("# 会话列表及会话盒子 — 50 条用例当前状态\n")
out.append(f"**统计**：共 50 条 ｜ ✅ 已通过 {len(passed)} ｜ ⏭️ 跳过 {len(skipped)} ｜ ⚠️ 待修复 {len(todo)}\n")

out.append(f"## ✅ 已通过 ({len(passed)})\n")
out.append("| Row | 测试点 | 测试文件 | 通过日期 |")
out.append("|---|---|---|---|")
for c in sorted(passed, key=lambda x: x['row']):
    out.append(f"| {c['row']} | {c['测试点']} | `{c['test_file']}` | {c['date']} |")

out.append(f"\n## ⏭️ 跳过 ({len(skipped)})\n")
out.append("| Row | 测试点 | 测试文件 | 跳过原因 |")
out.append("|---|---|---|---|")
for c in sorted(skipped, key=lambda x: x['row']):
    rm = re.search(r'跳过 \d{4}-\d{2}-\d{2} - (.*?) - 已实现', c['progress'])
    reason = rm.group(1) if rm else "?"
    out.append(f"| {c['row']} | {c['测试点']} | `{c['test_file']}` | {reason} |")

out.append(f"\n## ⚠️ 待修复 ({len(todo)})\n")
out.append("| Row | 测试点 | 测试文件 | 之前失败根因（来自最近回归） |")
out.append("|---|---|---|---|")

# 从 fail_cluster 拿失败根因
import json, datetime
today = datetime.date.today().isoformat()
fails = json.load(open(rf'E:\UIAutotest\logs\regression_failures_{today}.json', encoding='utf-8'))
row_to_root = {}
for f in fails:
    stdout = f.get('stdout_tail', '') or ''
    rm = re.search(r'(?:ERROR|FAILED) [^\s]+ - (.+)', stdout)
    if rm:
        # 只取第一个 row 的根因（同一 row 可能多次出现取最早）
        row_to_root.setdefault(f['row'], rm.group(1).strip()[:120])

for c in sorted(todo, key=lambda x: x['row']):
    root = row_to_root.get(c['row'], '(未跑过 / 无记录)')
    out.append(f"| {c['row']} | {c['测试点']} | `{c['test_file']}` | {root} |")

# 失败聚类
out.append(f"\n## 待修复用例的根因分类\n")
from collections import Counter
roots = [row_to_root.get(c['row'], '(无记录)') for c in todo]
# 简化签名
def normalize(r):
    r = re.sub(r"'[^']*'", "'X'", r)
    r = re.sub(r"\d+", "N", r)
    return r[:120]
sig_count = Counter(normalize(r) for r in roots)
out.append("| 错误签名 | 涉及条数 |")
out.append("|---|---|")
for sig, n in sig_count.most_common():
    out.append(f"| {sig} | {n} |")

Path(rf'E:\UIAutotest\logs\status_report_{today}.md').write_text("\n".join(out), encoding='utf-8')
print(f'写入 logs/status_report_{today}.md')
print(f'\n总计 {len(cases)} 条 = ✅ {len(passed)} + ⏭️ {len(skipped)} + ⚠️ {len(todo)}')
